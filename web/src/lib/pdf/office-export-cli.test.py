"""Real converter/OCR checks; run inside the isolated PDF runtime image."""
import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from docx import Document
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont

spec = importlib.util.spec_from_file_location("fixture", Path(__file__).with_name("office-export-excel.test.py"))
fixture = importlib.util.module_from_spec(spec)
spec.loader.exec_module(fixture)


class ExportCliTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)

    def convert(self, source, extension):
        output = self.directory / ("result." + extension)
        run = subprocess.run([sys.executable, str(Path(__file__).with_name("office-export.py")),
                              str(source), str(output), extension], capture_output=True, text=True, timeout=120)
        events = [json.loads(line) for line in run.stdout.splitlines()]
        self.assertEqual(run.returncode, 0, events)
        self.assertEqual(run.stderr, "")
        return output, events

    def test_digital_pdf_has_real_editable_word_and_excel(self):
        source = self.directory / "source.pdf"
        fixture.make_pdf(source, [(40, 750, "Relatorio de teste"), (40, 700, "123,45")])
        word, events = self.convert(source, "docx")
        self.assertIn("Relatorio de teste", " ".join(p.text for p in Document(word).paragraphs))
        self.assertEqual(events[1]["ocrPages"], 0)
        excel, _ = self.convert(source, "xlsx")
        book = load_workbook(excel)
        self.addCleanup(book.close)
        self.assertIn(123.45, [c.value for row in book.active for c in row])

    def test_scanned_pdf_ocr_produces_editable_text(self):
        source = self.directory / "scan.pdf"
        image = Image.new("RGB", (1654, 2339), "white")
        font = ImageFont.truetype("/usr/share/fonts/dejavu/DejaVuSans.ttf", 46)
        draw = ImageDraw.Draw(image)
        draw.text((120, 160), "RELATORIO DE CONVERSAO", fill="black", font=font)
        draw.text((120, 240), "Documento digitalizado para teste", fill="black", font=font)
        image.save(source, "PDF", resolution=200)
        image.close()
        output, events = self.convert(source, "docx")
        text = " ".join(p.text for p in Document(output).paragraphs)
        self.assertIn("RELATORIO", text.upper())
        self.assertIn("digitalizado", text.lower())
        self.assertEqual(events[1]["ocrPages"], 1)
        excel, events = self.convert(source, "xlsx")
        book = load_workbook(excel)
        self.addCleanup(book.close)
        text = " ".join(str(c.value or "") for row in book.active for c in row)
        self.assertIn("digitalizado", text.lower())
        self.assertEqual(events[1]["ocrPages"], 1)

    def test_page_limit_is_explicit_before_export(self):
        spec = importlib.util.spec_from_file_location("word_fixture", Path(__file__).with_name("office-export-word.test.py"))
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        source = self.directory / "many.pdf"
        module.make_pdf(source, [{"texts": [(40, 700, "Teste")]}] * 101)
        run = subprocess.run([sys.executable, str(Path(__file__).with_name("office-export.py")),
                              str(source), str(self.directory / "out.docx"), "docx"], capture_output=True, text=True)
        self.assertEqual(run.returncode, 1)
        self.assertEqual(json.loads(run.stdout.splitlines()[-1])["error"], "PDF_OFFICE_PAGE_LIMIT")

    def test_invalid_pdf_returns_code_without_payload(self):
        source = self.directory / "invalid.pdf"
        source.write_bytes(b"PRIVATE_DOCUMENT_CONTENT")
        run = subprocess.run([sys.executable, str(Path(__file__).with_name("office-export.py")),
                              str(source), str(self.directory / "out.docx"), "docx"], capture_output=True, text=True)
        self.assertNotEqual(run.returncode, 0)
        self.assertNotIn("PRIVATE_DOCUMENT_CONTENT", run.stdout + run.stderr)
        self.assertEqual(json.loads(run.stdout.splitlines()[-1])["error"], "PDF_OFFICE_CONVERSION_FAILED")


if __name__ == "__main__":
    unittest.main()
