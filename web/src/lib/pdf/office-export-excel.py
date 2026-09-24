"""Reconstruct editable worksheets without interpreting PDF text as formulas."""

import re
from bisect import bisect_left
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _value(text):
    text = str(text or "").strip()
    # Bare integers can be identifiers. Only explicit decimal/currency notation
    # is sufficiently unambiguous to promote to a spreadsheet number.
    match = re.fullmatch(r"(R\$\s*)?(-?(?:[1-9]\d{0,2}(?:\.\d{3})+|[1-9]\d*|0),\d{1,6})", text)
    if match:
        number = match.group(2).replace(".", "").replace(",", ".")
        if len(number.replace("-", "").replace(".", "")) <= 15:
            try:
                return Decimal(number), '"R$" #,##0.00' if match.group(1) else "0." + "0" * len(number.split(".")[1])
            except InvalidOperation:
                pass
    return text, None


def _write(sheet, row, column, text, header=False):
    value, number_format = _value(text)
    cell = sheet.cell(row, column, value)
    if isinstance(value, str):
        # Assignment normally treats '=...' as formula and '#N/A' as an error.
        cell.data_type = "s"
    if number_format:
        cell.number_format = number_format
    cell.alignment = Alignment(horizontal="right" if number_format else "left", vertical="top", wrap_text=True)
    if header:
        cell.font = Font(bold=True, color="17365D")
        cell.fill = PatternFill("solid", fgColor="E8EFF7")
    return cell


def _inside(word, bbox):
    x = (word["x0"] + word["x1"]) / 2
    y = (word["top"] + word["bottom"]) / 2
    return bbox[0] - 1 <= x <= bbox[2] + 1 and bbox[1] - 1 <= y <= bbox[3] + 1


def _rows(words):
    rows = []
    for word in sorted(words, key=lambda item: (item["top"], item["x0"])):
        if not rows or abs(word["top"] - rows[-1][0]["top"]) > 3:
            rows.append([word])
        else:
            rows[-1].append(word)
    return [sorted(row, key=lambda item: item["x0"]) for row in rows]


def _confident_borderless(table, words):
    data = [row for row in table.extract() if any(cell and cell.strip() for cell in row)]
    if len(data) < 3 or len(data[0]) < 2 or len(data[0]) > 30:
        return False
    columns = len(data[0])
    if sum(sum(bool(cell and cell.strip()) for cell in row) >= columns * .7 for row in data) < len(data) * .8:
        return False
    # Text strategy can split an ordinary paragraph at every aligned word.
    # Require genuinely separated fields on most physical lines.
    lines = _rows([word for word in words if _inside(word, table.bbox)])
    separated = sum(any(right["x0"] - left["x1"] >= 14 for left, right in zip(line, line[1:])) for line in lines)
    return len(lines) >= 3 and separated >= len(lines) * .7


def _tables(page, words):
    ruled = [table for table in page.find_tables() if len(table.rows) >= 2 and len(table.columns) >= 2]
    candidates = page.find_tables({
        "vertical_strategy": "text", "horizontal_strategy": "text",
        "min_words_vertical": 3, "min_words_horizontal": 1,
        "text_x_tolerance": 2, "text_y_tolerance": 3,
    })
    for table in candidates:
        # Keep ruled extraction authoritative; overlapping text tables duplicate it.
        overlaps = any(not (table.bbox[2] <= other.bbox[0] or table.bbox[0] >= other.bbox[2]
                                or table.bbox[3] <= other.bbox[1] or table.bbox[1] >= other.bbox[3]) for other in ruled)
        if not overlaps and _confident_borderless(table, words):
            ruled.append(table)
    return sorted(ruled, key=lambda table: (table.bbox[1], table.bbox[0]))


def _write_table(sheet, start, table, adjacent=()):
    data = table.extract(x_tolerance=2, y_tolerance=3)
    boundaries = sorted(set(round(cell[0], 2) for cell in table.cells) | set(round(cell[2], 2) for cell in table.cells))
    row_number = start
    first = True
    placed = set()
    source_rows = list(zip(table.rows, data))
    row_tops = [source_row.bbox[1] for source_row, _ in source_rows]
    row_bottom = max(cell[3] for cell in table.cells)
    retained = []
    for index, (source_row, cells) in enumerate(source_rows):
        top = row_tops[index]
        bottom = row_tops[index + 1] if index + 1 < len(row_tops) else row_bottom
        # A spanning cell makes source_row.bbox taller than its logical row.
        # Use consecutive row origins so side annotations stay on their record.
        row_words = [word for word in adjacent if top <= (word["top"]+word["bottom"])/2 < bottom]
        covered = any(cell[1] < top and cell[3] > top for cell in table.cells)
        if not any(value and value.strip() for value in cells) and not row_words and not covered:
            continue
        retained.append((source_row, cells, row_words))
    retained_tops = [row.bbox[1] for row, _, _ in retained]
    for source_row, cells, row_words in retained:
        # Coordinates, rather than None values, identify genuine spanning cells.
        for rectangle, text in zip(source_row.cells, cells):
            if rectangle is None:
                continue
            left = min(range(len(boundaries)), key=lambda index: abs(boundaries[index] - rectangle[0]))
            right = min(range(len(boundaries)), key=lambda index: abs(boundaries[index] - rectangle[2]))
            column = left + 1
            _write(sheet, row_number, column, text, header=first)
            last_row = start + bisect_left(retained_tops, rectangle[3] - .01) - 1
            if right > left + 1 or last_row > row_number:
                sheet.merge_cells(start_row=row_number, end_row=max(row_number, last_row), start_column=column, end_column=right)
        fields = {}
        for word in row_words:
            center = (word["x0"]+word["x1"])/2
            column = next((index+1 for index in range(len(boundaries)-1) if boundaries[index] <= center < boundaries[index+1]), len(boundaries))
            fields.setdefault(column, []).append(word)
        for column, words in fields.items():
            while isinstance(sheet.cell(row_number, column), MergedCell) or sheet.cell(row_number, column).value is not None:
                column = max(column + 1, len(boundaries))
            _write(sheet, row_number, column, "\n".join(" ".join(w["text"] for w in line) for line in _rows(words)))
            placed.update(id(word) for word in words)
        first = False
        row_number += 1
    # Irregular table geometry must never make unmatched annotations disappear.
    return _write_text(sheet, row_number, [word for word in adjacent if id(word) not in placed])


def _write_text(sheet, start, words):
    for line in _rows(words):
        # Keep sentences intact; large whitespace separates real column fields.
        fields = [[]]
        previous = None
        for word in line:
            if previous is not None and word["x0"] - previous["x1"] >= 24:
                fields.append([])
            fields[-1].append(word["text"])
            previous = word
        for column, field in enumerate(fields, 1):
            _write(sheet, start, column, " ".join(field))
        start += 1
    return start


def export_excel(pdf, output_path):
    workbook = Workbook()
    workbook.remove(workbook.active)
    table_count = 0
    for page_number, original in enumerate(pdf.pages, 1):
        page = original
        try:
            if original.cropbox and original.cropbox != original.bbox:
                page = original.crop(original.cropbox, strict=False)
            sheet = workbook.create_sheet(f"Página {page_number}")
            words = page.extract_words(x_tolerance=2, y_tolerance=3)
            tables = _tables(page, words)
            table_count += len(tables)
            # A detected table may have holes (row labels beside partial borders).
            # Exclude only occupied cells, before grouping characters into words.
            cells = [cell for table in tables for cell in table.cells]
            outside_page = page.filter(lambda obj: obj.get("object_type") != "char" or not any(
                cell[0] <= (obj["x0"]+obj["x1"])/2 <= cell[2]
                and cell[1] <= (obj["top"]+obj["bottom"])/2 <= cell[3] for cell in cells))
            outside = outside_page.extract_words(x_tolerance=2, y_tolerance=3)
            row = 1
            first_table_row = None
            for table in tables:
                before = [word for word in outside if word["top"] < table.bbox[1]]
                outside = [word for word in outside if word["top"] >= table.bbox[1]]
                row = _write_text(sheet, row, before)
                if first_table_row is None:
                    first_table_row = row
                # Preserve row labels in table holes and annotations beside their
                # corresponding records, rather than appending orphaned rows.
                adjacent = [word for word in outside if table.bbox[1] <= (word["top"]+word["bottom"])/2 < table.bbox[3]]
                outside = [word for word in outside if word not in adjacent]
                row = _write_table(sheet, row, table, adjacent) + 1
            _write_text(sheet, row, outside)
            if not words and not tables:
                _write(sheet, 1, 1, "Página sem texto reconhecível.")
            for column in sheet.columns:
                populated = [cell for cell in column if cell.value is not None]
                if populated:
                    width = max(max(len(line) for line in str(cell.value).splitlines() or [""]) for cell in populated)
                    sheet.column_dimensions[get_column_letter(column[0].column)].width = min(70, max(12, width + 2))
            sheet.freeze_panes = f"A{first_table_row + 1}" if first_table_row is not None else "A2"
            # Repeating one table's header on other tables mislabels their data.
            if len(tables) == 1 and first_table_row is not None:
                sheet.print_title_rows = f"{first_table_row}:{first_table_row}"
            sheet.sheet_view.showGridLines = True
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 0
            sheet.page_setup.orientation = "landscape" if page.width > page.height else "portrait"
            sheet.print_area = sheet.calculate_dimension()
        finally:
            page.close()
            if page is not original:
                original.close()
    if not workbook.worksheets:
        workbook.create_sheet("Documento")
    workbook.save(output_path)
    workbook.close()
    return {"pages": len(pdf.pages), "tables": table_count}
