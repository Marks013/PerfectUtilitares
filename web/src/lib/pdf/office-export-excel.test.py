import importlib.util
import tempfile
import unittest
from types import SimpleNamespace
from pathlib import Path

import pdfplumber
from openpyxl import Workbook, load_workbook

spec = importlib.util.spec_from_file_location("excel_export", Path(__file__).with_name("office-export-excel.py"))
engine = importlib.util.module_from_spec(spec)
spec.loader.exec_module(engine)


def make_pdf(path, texts, lines=(), page_options=""):
    commands = []
    for x, y, text in texts:
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        commands.append(f"BT /F1 10 Tf {x} {y} Td ({escaped}) Tj ET")
    for x1, y1, x2, y2 in lines:
        commands.append(f"{x1} {y1} m {x2} {y2} l S")
    stream = "\n".join(commands).encode("latin1")
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
               f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 600 800] {page_options} /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>".encode(),
               b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
               f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream"]
    content = b"%PDF-1.4\n"
    offsets = [0]
    for number, obj in enumerate(objects, 1):
        offsets.append(len(content))
        content += f"{number} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(content)
    content += f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode()
    content += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:])
    content += f"trailer << /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    path.write_bytes(content)


class ExcelExportTests(unittest.TestCase):
    def convert(self, texts, lines=(), page_options=""):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        path = Path(temporary.name)
        make_pdf(path / "source.pdf", texts, lines, page_options)
        with pdfplumber.open(path / "source.pdf") as pdf:
            metrics = engine.export_excel(pdf, path / "result.xlsx")
        book = load_workbook(path / "result.xlsx")
        self.addCleanup(book.close)
        return book.active, metrics

    def test_ruled_table_numbers_identifiers_and_outside_text(self):
        texts = [(40, 770, "Resumo"), (45, 725, "Codigo"), (205, 725, "Valor"),
                 (45, 695, "00123"), (205, 695, "R$ 1.234,56"),
                 (45, 665, "12345678901"), (205, 665, "-12,50"), (40, 610, "Observacao final")]
        lines = [(40, y, 400, y) for y in [740, 710, 680, 650]] + [(x, 650, x, 740) for x in [40, 200, 400]]
        sheet, metrics = self.convert(texts, lines)
        self.assertEqual(metrics, {"pages": 1, "tables": 1})
        self.assertEqual(sheet["A1"].value, "Resumo")
        self.assertEqual(sheet["A3"].value, "00123")
        self.assertEqual(sheet["B3"].value, 1234.56)
        self.assertEqual(sheet["A4"].value, "12345678901")
        self.assertEqual(sheet["B4"].value, -12.5)
        self.assertIn("Observacao final", [cell.value for row in sheet for cell in row])

    def test_formula_text_never_executes_and_prose_is_preserved(self):
        sheet, metrics = self.convert([(40, 750, '=HYPERLINK("https://example.org")'),
                                      (40, 725, "Texto simples preservado"), (40, 700, "#N/A")])
        self.assertEqual(metrics["tables"], 0)
        self.assertEqual(sheet["A1"].data_type, "s")
        self.assertEqual(sheet["A3"].data_type, "s")
        self.assertEqual(sheet["A2"].value, "Texto simples preservado")

    def test_borderless_table_has_real_columns(self):
        texts = [(40, y, left) for y, left in [(750, "Produto"), (720, "Caneta"), (690, "Papel"), (660, "Livro")]]
        texts += [(250, y, right) for y, right in [(750, "Valor"), (720, "12,50"), (690, "25,00"), (660, "30,00")]]
        sheet, metrics = self.convert(texts)
        self.assertEqual(metrics["tables"], 1)
        self.assertEqual(sheet["A2"].value, "Caneta")
        self.assertEqual(sheet["B2"].value, 12.5)

    def test_spanning_header_is_merged(self):
        texts = [(45, 725, "Cabecalho"), (45, 695, "Item"), (205, 695, "Valor"), (45, 665, "Papel"), (205, 665, "12,50")]
        lines = [(40, y, 400, y) for y in [740, 710, 680, 650]] + [(x, 650, x, 740) for x in [40, 400]] + [(200, 650, 200, 710)]
        sheet, metrics = self.convert(texts, lines)
        self.assertEqual(metrics["tables"], 1)
        self.assertIn("A1:B1", [str(item) for item in sheet.merged_cells.ranges])
        self.assertEqual(sheet["B3"].value, 12.5)

    def test_ambiguous_numbers_stay_text(self):
        for text in ["000123", "012,34", "1234567890123456,00", "123.456.789-01", "=1+1", "12345", "1.234"]:
            self.assertEqual(engine._value(text), (text, None))

    def test_cropbox_excludes_hidden_text(self):
        sheet, _ = self.convert([(40, 750, "Visible"), (350, 750, "Hidden")], page_options="/CropBox [0 0 300 800]")
        values = [str(cell.value) for row in sheet for cell in row if cell.value]
        self.assertEqual(values, ["Visible"])

    def test_rotated_page_keeps_text(self):
        sheet, _ = self.convert([(40, 750, "Rotation")], page_options="/Rotate 90")
        self.assertIn("Rotation", [cell.value for row in sheet for cell in row])

    def test_table_holes_preserve_row_identifiers(self):
        lines = [(80, y, 360, y) for y in [740, 725, 710, 695]]
        lines += [(x, 695, x, 740) for x in [80, 230, 360]]
        lines += [(40, 740, 80, 740), (40, 725, 80, 725), (40, 725, 40, 740)]
        sheet, _ = self.convert([
            (45, 730, "ID"), (85, 730, "Item"), (235, 730, "Value"),
            (45, 715, "001"), (85, 715, "Alpha"), (235, 715, "10,50"),
            (45, 700, "002"), (85, 700, "Beta"), (235, 700, "20,50"),
        ], lines)
        text = " ".join(str(cell.value) for row in sheet for cell in row if cell.value is not None)
        for value in ["001", "002", "Alpha", "Beta"]:
            self.assertEqual(text.count(value), 1)
        self.assertEqual(sheet['A2'].value, '001')
        self.assertEqual(sheet['B2'].value, 'Alpha')
        self.assertEqual(sheet['A3'].value, '002')
        self.assertEqual(sheet['B3'].value, 'Beta')

    def test_irregular_table_keeps_colliding_and_unmatched_annotations(self):
        sheet = Workbook().active
        table = SimpleNamespace(cells=[(10, 10, 90, 30)],
            rows=[SimpleNamespace(bbox=(10, 10, 90, 30), cells=[(10, 10, 90, 30)])],
            extract=lambda **_: [['Occupied']])
        words = [dict(x0=20, x1=40, top=12, bottom=22, text='Annotation'),
                 dict(x0=20, x1=40, top=35, bottom=45, text='Unmatched')]
        engine._write_table(sheet, 1, table, words)
        self.assertEqual(sheet['A1'].value, 'Occupied')
        self.assertEqual(sheet['B1'].value, 'Annotation')
        self.assertEqual(sheet['A2'].value, 'Unmatched')


if __name__ == "__main__":
    unittest.main()
