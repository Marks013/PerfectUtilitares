"""Reconstruct editable Word blocks from PDF geometry, without remote services."""
import io
import re
import subprocess
from functools import lru_cache
from statistics import median

import pypdfium2 as pdfium
from PIL import ImageFont

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


def clean(text):
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)


def inside(obj, box):
    x = (obj["x0"] + obj["x1"]) / 2
    y = (obj["top"] + obj["bottom"]) / 2
    return box[0] <= x <= box[2] and box[1] <= y <= box[3]


@lru_cache(maxsize=64)
def measured_font(family, bold, italic):
    pattern = family + (":style=Bold Italic" if bold and italic else ":style=Bold" if bold else ":style=Italic" if italic else "")
    result = subprocess.run(["fc-match", "-f", "%{file}", "--", pattern], capture_output=True, text=True, timeout=5, check=True)
    return ImageFont.truetype(result.stdout.strip(), 1000)


def font_run(paragraph, text, chars, scale=1):
    run = paragraph.add_run(clean(text))
    if not chars:
        return run
    char = chars[0]
    font = re.sub(r"^[A-Z]{6}\+", "", char.get("fontname", "Arial"))
    lowered = font.lower()
    run.bold = any(v in lowered for v in ("bold", "black", "semibold"))
    run.italic = any(v in lowered for v in ("italic", "oblique"))
    family = re.sub(r"[-,](Bold|Italic|Oblique|Regular|Roman|MT).*", "", font, flags=re.I)
    if family.startswith("___WRD_EMBED_SUB_"):
        family = "Arial"
    run.font.name = {"Helvetica": "Arial", "Calibri": "Arial", "Times": "Times New Roman", "Courier": "Courier New"}.get(family, family)
    run.font.size = Pt(max(3, min(144, float(char.get("size", 11)) * scale)))
    if chars and text.strip():
        source_width = max(c["x1"] for c in chars)-min(c["x0"] for c in chars)
        natural_width = measured_font(run.font.name, bool(run.bold), bool(run.italic)).getlength(text.strip()) / 1000 * run.font.size.pt
        if natural_width > 0:
            stretch = OxmlElement("w:w")
            stretch.set(qn("w:val"), str(max(10, min(300, round(source_width*scale/natural_width*100)))))
            run._r.get_or_add_rPr().append(stretch)
    color = char.get("non_stroking_color")
    if isinstance(color, (int, float)):
        color = (color, color, color)
    if isinstance(color, (list, tuple)) and len(color) in (1, 3, 4):
        if len(color) == 1:
            color = color * 3
        elif len(color) == 4:
            c, m, y, k = color
            color = ((1-c)*(1-k), (1-m)*(1-k), (1-y)*(1-k))
        run.font.color.rgb = RGBColor(*(max(0, min(255, round(v * 255))) for v in color))
    return run


def text_lines(page, excluded):
    # Filter characters before word grouping. A word may cross a table boundary,
    # and a table's bounding rectangle may contain genuine holes.
    if excluded:
        page = page.filter(lambda char: char.get("object_type") != "char" or not any(inside(char, box) for box in excluded))
    words = page.extract_words(x_tolerance=2, y_tolerance=3, return_chars=True,
                               extra_attrs=["fontname", "size"])
    rows = []
    for word in sorted(words, key=lambda w: (round(w["top"] / 3), w["x0"])):
        match = next((row for row in reversed(rows[-3:])
                      if abs(row["top"] - word["top"]) <= max(3, word["size"] * .25)), None)
        if match is None:
            match = {"kind": "text", "top": word["top"], "bottom": word["bottom"],
                     "x0": word["x0"], "x1": word["x1"], "words": []}
            rows.append(match)
        match["words"].append(word)
        match["bottom"] = max(match["bottom"], word["bottom"])
        match["x0"] = min(match["x0"], word["x0"])
        match["x1"] = max(match["x1"], word["x1"])
    # A large horizontal gap is a column boundary, not a string of spaces.
    lines = []
    for row in rows:
        group = []
        for word in sorted(row["words"], key=lambda w: w["x0"]):
            if group and word["x0"] - group[-1]["x1"] > max(24, word["size"] * 3):
                lines.append({**row, "words": group, "x0": group[0]["x0"], "x1": group[-1]["x1"]})
                group = []
            group.append(word)
        if group:
            lines.append({**row, "words": group, "x0": group[0]["x0"], "x1": group[-1]["x1"]})
    return lines


def find_tables(page):
    tables = [t for t in page.find_tables() if len(t.rows) >= 2 and len(t.columns) >= 2]
    if tables:
        return tables
    # Only accept borderless grids with repeated numeric data. Prose columns
    # otherwise frequently masquerade as tables.
    candidates = page.find_tables({"vertical_strategy": "text", "horizontal_strategy": "text",
                                   "min_words_vertical": 3, "intersection_tolerance": 5})
    for table in candidates:
        rows = [r for r in table.extract() if any(r)]
        values = [v for r in rows for v in r if v and v.strip()]
        numeric = sum(bool(re.fullmatch(r"(?:R\$\s*)?[-+]?\d[\d., ]*%?", v.strip())) for v in values)
        if len(rows) >= 3 and len(table.columns) >= 2 and numeric >= max(3, len(rows)-1) and len(values) >= len(rows) * 2:
            table._office_borderless = True
            tables.append(table)
    return tables


def put_table(parent, event, page, scale, page_image=None):
    source = event["table"]
    rows = source.rows
    xs = sorted({round(c[0], 2) for c in source.cells} | {round(c[2], 2) for c in source.cells})
    ys = sorted({round(c[1], 2) for c in source.cells} | {round(c[3], 2) for c in source.cells})
    table = parent.add_table(rows=len(ys)-1, cols=len(xs)-1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.style = "Normal Table" if getattr(source, "_office_borderless", False) else "Table Grid"
    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:w"), "0")
    indent.set(qn("w:type"), "dxa")
    table._tbl.tblPr.append(indent)
    margins = OxmlElement("w:tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        node = OxmlElement("w:" + side)
        node.set(qn("w:w"), "0")
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    table._tbl.tblPr.append(margins)
    for i, col in enumerate(table.columns):
        col.width = Pt((xs[i+1]-xs[i]) * scale)
    for row_index, row in enumerate(table.rows):
        row.height = Pt((ys[row_index+1]-ys[row_index]) * scale)
        row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
        for column_index, cell in enumerate(row.cells):
            cell.width = Pt((xs[column_index+1]-xs[column_index]) * scale)
            # Preserve narrow PDF rows rather than imposing Word's default
            # 11-point empty paragraphs on every blank cell.
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            cell.paragraphs[0].paragraph_format.line_spacing = 1
            cell.paragraphs[0].add_run().font.size = Pt(1)
    used = set()
    for box in source.cells:
        r0, r1 = ys.index(round(box[1], 2)), ys.index(round(box[3], 2))
        c0, c1 = xs.index(round(box[0], 2)), xs.index(round(box[2], 2))
        if (r0, c0) in used:
            continue
        cell = table.cell(r0, c0)
        if r1-r0 > 1 or c1-c0 > 1:
            cell = cell.merge(table.cell(r1-1, c1-1))
        # Merging cells concatenates their empty paragraphs; remove those before
        # filling the single PDF cell so they do not inflate the table height.
        for empty in list(cell.paragraphs)[1:]:
            if not empty.text:
                empty._p.getparent().remove(empty._p)
        used.update((r, c) for r in range(r0, r1) for c in range(c0, c1))
        backgrounds = [rect for rect in page.rects if rect.get("fill")
                       and rect["x0"] <= box[0]+1 and rect["x1"] >= box[2]-1
                       and rect["top"] <= box[1]+1 and rect["bottom"] >= box[3]-1]
        if backgrounds:
            background = min(backgrounds, key=lambda rect: rect["width"]*rect["height"])
            color = background.get("non_stroking_color")
            if isinstance(color, (tuple, list)) and len(color) == 3:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "".join(f"{max(0, min(255, round(v*255))):02X}" for v in color))
                cell._tc.get_or_add_tcPr().append(shading)
        cropped = page.filter(lambda obj: obj.get("object_type") != "char" or inside(obj, box))
        lines = text_lines(cropped, [])
        for index, line in enumerate(sorted(lines, key=lambda v: (v["top"], v["x0"]))):
            p = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = Pt(max(w["size"] for w in line["words"]) * scale)
            if abs((line["x0"]+line["x1"])-(box[0]+box[2])) < 4:
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif box[2]-line["x1"] < 3 and line["x0"]-box[0] > 5:
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for j, word in enumerate(line["words"]):
                font_run(p, (" " if j else "") + word["text"], word["chars"], scale)
        if page_image is not None:
            for image in page.images:
                if is_scanned_background(page, image) or not contains_box(box, (image["x0"], image["top"], image["x1"], image["bottom"])):
                    continue
                image_box = (max(box[0], image["x0"]), max(box[1], image["top"]),
                             min(box[2], image["x1"]), min(box[3], image["bottom"]))
                if image_box[2]-image_box[0] <= 3 or image_box[3]-image_box[1] <= 3:
                    continue
                stream = io.BytesIO()
                region = page_image.crop(tuple(round((v-page.bbox[i % 2])*2) for i, v in enumerate(image_box)))
                region.save(stream, format="PNG")
                region.close()
                stream.seek(0)
                paragraph = cell.paragraphs[0] if not lines else cell.add_paragraph()
                paragraph.paragraph_format.line_spacing = Pt((image_box[3]-image_box[1])*scale)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.add_run().add_picture(stream, width=Pt((image_box[2]-image_box[0])*scale), height=Pt((image_box[3]-image_box[1])*scale))
    if rows:
        header_chars = [c for c in page.chars if inside(c, (source.bbox[0], ys[0], source.bbox[2], ys[1]))]
        if header_chars and sum("bold" in c.get("fontname", "").lower() for c in header_chars) > len(header_chars)/2:
            repeat = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
    return table


def position_paragraph(paragraph, box, page, scale):
    """Use editable Word frames for content beside tables or over artwork."""
    frame = OxmlElement("w:framePr")
    values = {"hAnchor": "page", "vAnchor": "page", "wrap": "none",
              "x": str(round((box[0]-page.bbox[0])*scale*20)),
              "y": str(round((box[1]-page.bbox[1])*scale*20)),
              "w": str(round((box[2]-box[0])*scale*20)),
              "hRule": "auto"}
    for name, value in values.items():
        frame.set(qn("w:" + name), value)
    paragraph._p.get_or_add_pPr().append(frame)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1


def positioned_text(document, lines, page, scale):
    for line in lines:
        p = document.add_paragraph()
        # Small allowance absorbs font substitution without moving the frame.
        width = min(page.bbox[2], line["x1"] + max(12, (line["x1"]-line["x0"])*.08))
        position_paragraph(p, (line["x0"], line["top"], width, line["bottom"]), page, scale)
        p.paragraph_format.line_spacing = Pt(max(w["size"] for w in line["words"]) * scale)
        for index, word in enumerate(line["words"]):
            font_run(p, (" " if index else "") + word["text"], word["chars"], scale)


def artwork_image(page):
    """Render graphics without native text, for backgrounds and table images."""
    stream = page.pdf.stream
    offset = stream.tell()
    stream.seek(0)
    source = pdfium.PdfDocument(stream)
    try:
        pdf_page = source[page.page_number-1]
        try:
            for obj in pdf_page.get_objects(filter=[pdfium.raw.FPDF_PAGEOBJ_TEXT]):
                if not pdfium.raw.FPDFPageObj_SetIsActive(obj, False):
                    raise RuntimeError("Cannot separate artwork from text")
            bitmap = pdf_page.render(scale=2, draw_annots=False)
            try:
                image = bitmap.to_pil().copy()
            finally:
                bitmap.close()
        finally:
            pdf_page.close()
    finally:
        source.close()
        stream.seek(offset)
    return image


def artwork_background(document, page, scale):
    """Render artwork alone, retaining native PDF text as editable Word frames."""
    image = artwork_image(page)
    data = io.BytesIO()
    try:
        image.save(data, format="PNG")
    finally:
        image.close()
    data.seek(0)
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.line_spacing = Pt(1)
    inline = paragraph.add_run().add_picture(data, width=Pt(page.width*scale), height=Pt(page.height*scale))._inline
    anchor = OxmlElement("wp:anchor")
    for key, value in {"distT":"0", "distB":"0", "distL":"0", "distR":"0", "simplePos":"0", "relativeHeight":"0", "behindDoc":"1", "locked":"0", "layoutInCell":"1", "allowOverlap":"1"}.items():
        anchor.set(key, value)
    simple = OxmlElement("wp:simplePos")
    simple.set("x", "0")
    simple.set("y", "0")
    anchor.append(simple)
    for axis in ("H", "V"):
        pos = OxmlElement("wp:position" + axis)
        pos.set("relativeFrom", "page")
        value = OxmlElement("wp:posOffset")
        value.text = "0"
        pos.append(value)
        anchor.append(pos)
    anchor.append(inline.extent)
    anchor.append(OxmlElement("wp:wrapNone"))
    for child in list(inline):
        anchor.append(child)
    inline.getparent().replace(inline, anchor)


def contains_box(outer, inner):
    return outer[0] <= inner[0] and outer[1] <= inner[1] and outer[2] >= inner[2] and outer[3] >= inner[3]


def is_scanned_background(page, image):
    return (image["x1"]-image["x0"])*(image["bottom"]-image["top"]) > page.width*page.height*.7 and bool(page.chars)


def graphics(page, tables):
    boxes = []
    for image in page.images:
        box = (image["x0"], image["top"], image["x1"], image["bottom"])
        if is_scanned_background(page, image):
            # Scanned background is replaced by editable OCR text, not duplicated.
            continue
        if box[2]-box[0] > 3 and box[3]-box[1] > 3:
            boxes.append(box)
    # Preserve vector artwork as image regions; table rules remain editable.
    for curve in page.curves:
        if curve["width"] > 8 and curve["height"] > 8 and not any(inside(curve, t.bbox) for t in tables):
            boxes.append((curve["x0"], curve["top"], curve["x1"], curve["bottom"]))
    merged = []
    for box in boxes:
        matches = [b for b in merged if not (box[2] < b[0]-2 or box[0] > b[2]+2 or box[3] < b[1]-2 or box[1] > b[3]+2)]
        for b in matches:
            merged.remove(b)
            box = (min(box[0], b[0]), min(box[1], b[1]), max(box[2], b[2]), max(box[3], b[3]))
        merged.append(box)
    return merged


def emit(parent, events, page, left, right, scale, page_image, cell_image=None):
    previous = None
    paragraph = None
    for event in sorted(events, key=lambda e: (round(e["top"] / 3), e["x0"])):
        gap = max(0, event["top"] - (previous["bottom"] if previous else page.bbox[1]))
        if event["kind"] == "table":
            if previous and gap > 0:
                spacer = parent.add_paragraph()
                spacer.paragraph_format.space_after = Pt(0)
                spacer.paragraph_format.line_spacing = Pt(max(1, gap * scale))
                spacer.add_run().font.size = Pt(1)
            put_table(parent, event, page, scale, cell_image)
            paragraph = None
        elif event["kind"] == "image":
            p = parent.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(gap * scale if previous else 0)
            p.paragraph_format.left_indent = Pt(max(0, event["x0"]-left)*scale)
            stream = io.BytesIO()
            page_image.crop((round((event["x0"]-page.bbox[0])*2), round((event["top"]-page.bbox[1])*2),
                             round((event["x1"]-page.bbox[0])*2), round((event["bottom"]-page.bbox[1])*2))).save(stream, format="PNG")
            stream.seek(0)
            p.add_run().add_picture(stream, width=Pt((event["x1"]-event["x0"])*scale))
            paragraph = None
        else:
            size = median(w["size"] for w in event["words"])
            join = (paragraph is not None and previous and previous["kind"] == "text"
                    and gap <= size*.6 and abs(event["x0"]-previous["x0"]) < size
                    and abs(size-median(w["size"] for w in previous["words"])) < 1
                    and previous["x1"] > right-size*4)
            if not join:
                paragraph = parent.add_paragraph()
                fmt = paragraph.paragraph_format
                fmt.space_before = Pt(gap*scale if previous else 0)
                fmt.space_after = Pt(0)
                fmt.line_spacing = 1
                fmt.widow_control = False
                fmt.left_indent = Pt(max(0, event["x0"]-left)*scale)
                if abs((event["x0"]+event["x1"])/2-(left+right)/2) < 5 and event["x1"]-event["x0"] < (right-left)*.8:
                    fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    fmt.left_indent = Pt(0)
            for index, word in enumerate(event["words"]):
                font_run(paragraph, (" " if index or join else "") + word["text"], word["chars"], scale)
        previous = event


def emit_columns(document, events, page, left, middle, right, scale, page_image, cell_image=None):
    if not events:
        return
    left_events = [e for e in events if e["x1"] <= middle]
    right_events = [e for e in events if e["x0"] >= middle]
    if not left_events or not right_events:
        emit(document, events, page, left, right, scale, page_image, cell_image)
        return
    columns = document.add_table(rows=1, cols=2)
    columns.autofit = False
    for cell, items, a, b in [(columns.cell(0, 0), left_events, left, middle),
                              (columns.cell(0, 1), right_events, middle, right)]:
        cell.width = Pt((b-a)*scale)
        emit(cell, items, page, a, b, scale, page_image, cell_image)
        first = cell.paragraphs[0]
        if not first.text:
            first._element.getparent().remove(first._element)


def export_word(pdf, output_path):
    document = Document()
    document.core_properties.author = "PerfectUtilitares"
    document.core_properties.title = "Documento convertido"
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(0)
    count = 0
    for index, original in enumerate(pdf.pages):
        page = original.crop(original.cropbox, strict=False) if original.cropbox else original
        section = document.sections[0] if index == 0 else document.add_section(WD_SECTION_START.NEW_PAGE)
        # PDF headers/footers remain body content. Creating empty Word headers
        # would reserve extra space and displace content near the page edge.
        section.header_distance = Pt(0)
        section.footer_distance = Pt(0)
        scale = min(1, 1500 / max(page.width, page.height))
        section.page_width = Pt(page.width * scale)
        section.page_height = Pt(page.height * scale)
        tables = find_tables(page)
        boxes = graphics(page, tables)
        cell_images = [box for box in boxes if any(
            contains_box(cell, box)
            for table in tables for cell in table.cells)]
        boxes = [box for box in boxes if box not in cell_images]
        # Decorative frames spanning a page must not swallow editable text.
        artwork = not tables and page.chars and any((b[2]-b[0])*(b[3]-b[1]) > page.width*page.height*.6 for b in boxes)
        if artwork:
            section.left_margin = section.right_margin = Pt(0)
            section.top_margin = section.bottom_margin = Pt(0)
            artwork_background(document, page, scale)
            positioned_text(document, text_lines(page, []), page, scale)
            if page is not original:
                page.close()
            original.close()
            continue
        lines = text_lines(page, [cell for t in tables for cell in t.cells] + boxes)
        events = lines + [{"kind": "table", "table": t, "x0": t.bbox[0], "top": t.bbox[1], "x1": t.bbox[2], "bottom": t.bbox[3]} for t in tables]
        events += [{"kind": "image", "x0": b[0], "top": b[1], "x1": b[2], "bottom": b[3]} for b in boxes]
        left = min((e["x0"] for e in events), default=page.bbox[0]+24)
        right = max((e["x1"] for e in events), default=page.bbox[2]-24)
        top = min((e["top"] for e in events), default=page.bbox[1]+24)
        section.left_margin = Pt(max(0, left-page.bbox[0])*scale)
        # Do not constrain flowing Word text to the exact width of PDF glyphs:
        # font substitutions otherwise wrap even single words onto extra lines.
        section.right_margin = Pt(min(36, max(0, page.bbox[2]-right))*scale)
        section.top_margin = Pt(max(0, top-page.bbox[1])*scale)
        section.bottom_margin = Pt(8)
        page_image = page.to_image(resolution=144).original if boxes else None
        cell_image = artwork_image(page) if cell_images else None
        # A stable central gutter becomes two independently editable columns.
        middle = (left+right)/2
        cross = [e for e in events if e["x0"] < middle < e["x1"]]
        left_events = [e for e in events if e["x1"] <= middle]
        right_events = [e for e in events if e["x0"] >= middle]
        if tables and any(any(e["top"] < t.bbox[3] and e["bottom"] > t.bbox[1] for t in tables) for e in lines):
            # Floating tables and text frames preserve parallel annotations.
            for t in tables:
                table = put_table(document, {"table": t}, page, scale, cell_image)
                position = OxmlElement("w:tblpPr")
                for key, value in {"horzAnchor":"page", "vertAnchor":"page", "tblpX":str(round((t.bbox[0]-page.bbox[0])*scale*20)), "tblpY":str(round((t.bbox[1]-page.bbox[1])*scale*20)), "leftFromText":"0", "rightFromText":"0", "topFromText":"0", "bottomFromText":"0"}.items():
                    position.set(qn("w:"+key), value)
                table._tbl.tblPr.append(position)
            positioned_text(document, lines, page, scale)
            emit(document, [e for e in events if e["kind"] == "image"], page, left, right, scale, page_image)
        elif len(left_events) >= 3 and len(right_events) >= 3 and len(cross) <= max(2, len(events) * .15):
            remaining = left_events + right_events
            for spanning in sorted(cross, key=lambda e: e["top"]):
                band = [e for e in remaining if e["top"] < spanning["top"]]
                remaining = [e for e in remaining if e["top"] >= spanning["top"]]
                emit_columns(document, band, page, left, middle, right, scale, page_image, cell_image)
                emit(document, [spanning], page, left, right, scale, page_image, cell_image)
            emit_columns(document, remaining, page, left, middle, right, scale, page_image, cell_image)
        else:
            emit(document, events, page, left, right, scale, page_image, cell_image)
        count += len(tables)
        if page_image:
            page_image.close()
        if cell_image:
            cell_image.close()
        if page is not original:
            page.close()
        original.close()
    document.save(output_path)
    return {"pages": len(pdf.pages), "tables": count}
