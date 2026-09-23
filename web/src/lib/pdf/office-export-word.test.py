"""Real PDF -> DOCX -> PDF regressions. Fixtures contain synthetic data only."""
import importlib.util
import re
import subprocess
import tempfile
import unittest
from pathlib import Path

import pdfplumber
from docx import Document
from docx.oxml.ns import qn

spec = importlib.util.spec_from_file_location("word_export", Path(__file__).with_name("office-export-word.py"))
engine = importlib.util.module_from_spec(spec)
spec.loader.exec_module(engine)


def make_pdf(path, pages):
    objects = [b"", b""]
    def add(value):
        objects.append(value)
        return len(objects)
    regular = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
    bold = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold /Encoding /WinAnsiEncoding >>")
    italic = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Oblique /Encoding /WinAnsiEncoding >>")
    image = add(b"<< /Type /XObject /Subtype /Image /Width 2 /Height 2 /ColorSpace /DeviceRGB /BitsPerComponent 8 /Length 12 >>\nstream\n" + bytes([255, 0, 0, 0, 255, 0, 0, 0, 255, 255, 255, 0]) + b"\nendstream")
    page_ids = []
    for page in pages:
        commands = []
        for x, y, text, *style in page.get("texts", []):
            font = style[0] if style else "F1"
            size = style[1] if len(style) > 1 else 10
            escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            commands.append(f"BT /{font} {size} Tf {x} {y} Td ({escaped}) Tj ET")
        for x1, y1, x2, y2 in page.get("lines", []):
            commands.append(f"{x1} {y1} m {x2} {y2} l S")
        if page.get("image"):
            commands.append("q 60 0 0 40 40 620 cm /Im1 Do Q")
        stream = "\n".join(commands).encode("latin1")
        content = add(f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream")
        page_ids.append(add(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 600 800] {page.get('options', '')} /Resources << /Font << /F1 {regular} 0 R /F2 {bold} 0 R /F3 {italic} 0 R >> /XObject << /Im1 {image} 0 R >> >> /Contents {content} 0 R >>".encode()))
    objects[0] = b"<< /Type /Catalog /Pages 2 0 R >>"
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(f'{index} 0 R' for index in page_ids)}] /Count {len(page_ids)} >>".encode()
    content = b"%PDF-1.4\n"
    offsets = []
    for number, obj in enumerate(objects, 1):
        offsets.append(len(content))
        content += f"{number} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(content)
    content += f"xref\n0 {len(objects)+1}\n0000000000 65535 f \n".encode()
    content += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets)
    content += f"trailer << /Size {len(objects)+1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    path.write_bytes(content)


class WordExportTests(unittest.TestCase):
    def convert(self, pages):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        path = Path(temporary.name)
        make_pdf(path / "source.pdf", pages)
        with pdfplumber.open(path / "source.pdf") as pdf:
            metrics = engine.export_word(pdf, path / "converted.docx")
        return Document(path / "converted.docx"), metrics, path

    def render(self, path):
        result = subprocess.run(["libreoffice", "-env:UserInstallation=" + (path / "profile").as_uri(), "--headless", "--convert-to", "pdf", "--outdir", str(path), str(path / "converted.docx")], capture_output=True, text=True, timeout=45)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertTrue((path / "converted.pdf").exists(), result.stdout + result.stderr)
        with pdfplumber.open(path / "converted.pdf") as pdf:
            return len(pdf.pages), "\n".join(page.extract_text() or "" for page in pdf.pages)

    def test_editable_fonts_and_image(self):
        document, metrics, path = self.convert([{"texts": [(40, 750, "Editable regular"), (40, 720, "Bold title", "F2", 16), (40, 685, "Italic text", "F3", 11)], "image": True}])
        runs = [run for paragraph in document.paragraphs for run in paragraph.runs]
        self.assertIn("Editable", " ".join(run.text for run in runs))
        self.assertTrue(any(run.bold and "Bold" in run.text for run in runs))
        self.assertTrue(any(run.italic and "Italic" in run.text for run in runs))
        self.assertEqual(len(document.inline_shapes), 1)
        self.assertEqual(metrics, {"pages": 1, "tables": 0})
        count, text = self.render(path)
        self.assertEqual(count, 1)
        self.assertEqual(text.count("Editable"), 1)
        self.assertEqual(text.count("Bold"), 1)

    def test_table_cells_and_merges(self):
        lines = [(40, y, 400, y) for y in [740, 710, 680, 650]] + [(x, 650, x, 740) for x in [40, 400]] + [(200, 650, 200, 710)]
        document, metrics, path = self.convert([{"texts": [(45, 725, "Summary", "F2"), (45, 695, "Item"), (205, 695, "Value"), (45, 665, "Paper"), (205, 665, "12,50")], "lines": lines}])
        self.assertEqual(metrics["tables"], 1)
        table = document.tables[0]
        self.assertEqual(table.cell(2, 0).text, "Paper")
        self.assertEqual(table.cell(2, 1).text, "12,50")
        self.assertIsNotNone(table.cell(0, 0)._tc.tcPr.find(qn("w:gridSpan")))
        count, text = self.render(path)
        self.assertEqual(count, 1)
        self.assertEqual(text.count("Summary"), 1)
        self.assertEqual(text.count("Paper"), 1)

    def test_two_columns_preserve_all_text(self):
        texts = [(40, y, left) for y, left in [(750, "Left Alpha"), (725, "Left Beta"), (700, "Left Gamma")]]
        texts += [(330, y, right) for y, right in [(750, "Right Delta"), (725, "Right Epsilon"), (700, "Right Zeta")]]
        document, metrics, path = self.convert([{"texts": texts}])
        self.assertEqual(metrics["tables"], 0)
        self.assertEqual(len(document.tables), 1)
        self.assertIn("Alpha", document.tables[0].cell(0, 0).text)
        self.assertIn("Delta", document.tables[0].cell(0, 1).text)
        count, text = self.render(path)
        self.assertEqual(count, 1)
        for token in ["Alpha", "Beta", "Gamma", "Delta", "Epsilon", "Zeta"]:
            self.assertEqual(text.count(token), 1)

    def test_cropbox_excludes_hidden_text(self):
        document, _, path = self.convert([{"texts": [(40, 750, "Visible"), (350, 750, "Hidden")], "options": "/CropBox [0 0 300 800]"}])
        self.assertEqual(document.sections[0].page_width.pt, 300)
        _, text = self.render(path)
        self.assertIn("Visible", text)
        self.assertNotIn("Hidden", text)

    def test_heading_above_two_columns_does_not_interleave_reading_order(self):
        texts = [(150, 780, "Document heading across columns", "F2", 18)]
        texts += [(40, y, f"Left {value}") for y, value in [(730, "Alpha"), (700, "Beta"), (670, "Gamma")]]
        texts += [(330, y, f"Right {value}") for y, value in [(730, "Delta"), (700, "Epsilon"), (670, "Zeta")]]
        document, _, path = self.convert([{"texts": texts}])
        self.assertIn("Document heading", " ".join(p.text for p in document.paragraphs))
        self.assertEqual(len(document.tables), 1)
        self.assertIn("Gamma", document.tables[0].cell(0, 0).text)
        self.assertNotIn("Delta", document.tables[0].cell(0, 0).text)
        count, text = self.render(path)
        self.assertEqual(count, 1)
        self.assertEqual(text.count("heading"), 1)

    def test_rotation_keeps_content(self):
        _, _, path = self.convert([{"texts": [(40, 750, "Rotation")], "options": "/Rotate 90"}])
        count, text = self.render(path)
        self.assertEqual(count, 1)
        self.assertIn("Rotation", re.sub(r"\s+", "", text))

    def test_portuguese_text_keeps_accents(self):
        document, _, path = self.convert([{"texts": [(40, 750, "Relatório: conversão, ação e informações.")]}])
        self.assertIn("conversão", " ".join(p.text for p in document.paragraphs))
        _, text = self.render(path)
        self.assertIn("informações", text)

    def test_two_physical_pages_do_not_overflow(self):
        pages = [{"texts": [(40, 750-i*14, f"Page{page} Line{i:02d} synthetic paragraph text for pagination check") for i in range(48)]} for page in [1, 2]]
        document, metrics, path = self.convert(pages)
        self.assertEqual(metrics["pages"], 2)
        self.assertEqual(len(document.sections), 2)
        count, text = self.render(path)
        self.assertEqual(count, 2)
        self.assertEqual(text.count("Page1"), 48)
        self.assertEqual(text.count("Page2"), 48)


if __name__ == "__main__":
    unittest.main()
